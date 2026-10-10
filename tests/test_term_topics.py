"""Term topics for the parent–teacher agent. Run with: pytest tests/test_term_topics.py"""

import json
import os

os.environ.setdefault("OPENAI_API_KEY", "test-key")
os.environ.setdefault("DATABASE_URL", "sqlite:///./dev.db")

import pytest
from openpyxl import load_workbook

from app.agents.tools import (
    ACADEMIC_TOOLS,
    COMMUNICATION_TOOLS,
    DAILY_UPDATE_TOOLS,
    INSIGHT_TOOLS,
    PARENT_TEACHER_TOOLS,
    get_term_topics_for_student,
)
from app.services.excel_teacher_store import write_demo_workbook

FA2_TITLES = ["Adding 3-digit numbers", "Number stories"]


@pytest.fixture
def workbook(tmp_path, monkeypatch):
    xlsx = tmp_path / "teacher_workspace.xlsx"
    write_demo_workbook(xlsx)
    monkeypatch.setattr("app.services.excel_teacher_store.settings.teacher_excel_path", str(xlsx))
    return xlsx


def _topics(student_id: str, subject: str, term: str = "") -> dict:
    return get_term_topics_for_student.invoke(
        {"student_id": student_id, "subject": subject, "term": term}
    )


def test_known_class_returns_topics_and_teacher(workbook):
    result = _topics("s-3-1", "Mathematics")
    assert result["subject"] == "Mathematics"
    assert result["term"] == "FA2"
    assert result["teacher"] == "Kavita Nair"
    assert [row["title"] for row in result["topics"]] == FA2_TITLES
    assert result["topics"][0]["kind"] == "topic"
    assert result["topics"][0]["status"] == "in_progress"
    assert result["topics"][0]["description"]
    assert result["topics"][1]["kind"] == "project"
    assert set(result.keys()) == {"subject", "term", "teacher", "topics"}
    for row in result["topics"]:
        assert set(row.keys()) == {"title", "kind", "status", "description"}


def test_empty_term_selects_latest_academic_term(workbook):
    result = _topics("s-3-1", "Mathematics", term="")
    assert result["term"] == "FA2"
    titles = [row["title"] for row in result["topics"]]
    assert titles == FA2_TITLES
    assert "Place value to 999" not in titles
    assert "Later label" not in titles


def test_subject_match_ignores_capitalisation(workbook):
    for subject in ("mathematics", "MATHEMATICS", "Mathematics"):
        result = _topics("s-3-1", subject)
        assert result["subject"] == "Mathematics"
        assert result["teacher"] == "Kavita Nair"
        assert [row["title"] for row in result["topics"]] == FA2_TITLES


def test_unknown_student(workbook):
    assert _topics("no-such-student", "Mathematics") == {"error": "student not found"}


def test_no_topics_includes_note(workbook):
    result = _topics("s-3-1", "History")
    assert result == {
        "topics": [],
        "note": "No History topics have been entered for this class yet.",
    }


def test_explicit_term_does_not_mix_other_terms(workbook):
    result = _topics("s-3-1", "Mathematics", term="FA1")
    assert result["term"] == "FA1"
    assert [row["title"] for row in result["topics"]] == ["Place value to 999"]

    missing = _topics("s-3-1", "Mathematics", term="SA2")
    assert missing == {
        "topics": [],
        "note": "No Mathematics topics have been entered for this class yet.",
    }


def test_other_classes_and_teachers_do_not_leak(workbook):
    grade_three = _topics("s-3-1", "Mathematics")
    titles = [row["title"] for row in grade_three["topics"]]
    assert "Grade 4 fractions" not in titles
    assert "Other teacher draft" not in titles
    assert grade_three["teacher"] == "Kavita Nair"

    grade_four = _topics("s3", "Mathematics")
    assert grade_four["teacher"] == "Arjun Rao"
    assert [row["title"] for row in grade_four["topics"]] == ["Grade 4 fractions"]
    assert "Adding 3-digit numbers" not in [row["title"] for row in grade_four["topics"]]


def test_result_omits_contact_details_and_internal_ids(workbook):
    wb = load_workbook(workbook)
    ws = wb["teachers"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    phone_col = len(headers) + 1
    ws.cell(1, phone_col, "phone")
    id_col = headers.index("id") + 1
    for row_idx in range(2, ws.max_row + 1):
        if ws.cell(row_idx, id_col).value == "t3":
            ws.cell(row_idx, phone_col, "9876543210")
    wb.save(workbook)
    wb.close()

    result = _topics("s-3-1", "Mathematics")
    blob = json.dumps(result)
    assert "kavita.nair@demo-school.edu" not in blob
    assert "9876543210" not in blob
    assert "demo-school" not in blob
    assert "t3" not in blob

    keys = set()

    def walk(value):
        if isinstance(value, dict):
            keys.update(value.keys())
            for item in value.values():
                walk(item)
        elif isinstance(value, list):
            for item in value:
                walk(item)

    walk(result)
    assert keys.isdisjoint({"id", "tenant_id", "teacher_id", "email", "phone", "phone_number"})


def test_tool_is_bound_only_to_parent_teacher_agent():
    import app.agents.academic_agent as academic_agent
    import app.agents.communication_agent as parent_agent
    import app.agents.daily_update_agent as daily_agent

    assert get_term_topics_for_student in PARENT_TEACHER_TOOLS
    assert get_term_topics_for_student in parent_agent.PARENT_TEACHER_TOOLS
    assert parent_agent.PARENT_TEACHER_TOOLS is PARENT_TEACHER_TOOLS
    for unrelated in (DAILY_UPDATE_TOOLS, ACADEMIC_TOOLS, COMMUNICATION_TOOLS, INSIGHT_TOOLS):
        assert get_term_topics_for_student not in unrelated
    assert get_term_topics_for_student not in academic_agent.ACADEMIC_TOOLS
    assert get_term_topics_for_student not in daily_agent.DAILY_UPDATE_TOOLS

    bound = parent_agent.llm
    bound_tools = getattr(bound, "kwargs", {}).get("tools") or []
    names = []
    for tool in bound_tools:
        if isinstance(tool, dict):
            names.append(tool.get("function", {}).get("name") or tool.get("name"))
        else:
            names.append(getattr(tool, "name", None))
    assert "get_term_topics_for_student" in names


def test_parent_teacher_prompt_uses_the_tool():
    from app.agents.communication_agent import SYSTEM_PROMPT

    text = " ".join(SYSTEM_PROMPT.split())
    assert "get_term_topics_for_student" in text
    assert "Never invent topics" in text
    assert "teacher will share the details" in text
