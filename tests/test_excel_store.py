"""ExcelStore reads, writes, reloads and errors, against a temp copy of the workbook."""

import pytest
from openpyxl import load_workbook

from app.services import excel_store
from app.services.excel_store import ExcelStore, WorkbookError, WorkbookLockedError, get_store


def first(rows):
    assert rows, "expected at least one row"
    return rows[0]


def test_cells_are_normalised(store):
    student = first(store.rows("students"))
    assert isinstance(student["active"], bool)
    assert len(student["date_of_birth"]) == 10 and student["date_of_birth"][4] == "-"
    assert isinstance(student["grade"], str)

    guardian = first(store.rows("guardians"))
    assert isinstance(guardian["whatsapp_number"], str)
    assert guardian["whatsapp_number"].isdigit()

    record = first(store.rows("assessment_records"))
    assert record["marks_total"] is None or isinstance(record["marks_total"], (int, float))


def test_get_and_find(store):
    student = first(store.rows("students"))
    assert store.get("students", student["id"]) == student
    assert store.get("students", "no-such-id") is None
    assert all(r["student_id"] == student["id"] for r in store.find("attendance", student_id=student["id"]))


def test_missing_sheet_is_an_error(store):
    with pytest.raises(WorkbookError):
        store.rows("no_such_sheet")


def test_get_store_is_one_shared_instance(workbook):
    assert get_store() is get_store()


def test_append_row_generates_id_and_persists(store, workbook):
    tenant_id = store.default_tenant_id()
    row = store.append_row("conversations", {"tenant_id": tenant_id, "guardian_id": "x", "channel": "app"})
    assert row["id"].startswith("conversations-")
    assert store.get("conversations", row["id"])["channel"] == "app"
    assert ExcelStore(workbook).get("conversations", row["id"]) is not None


def test_append_rejects_unknown_columns(store):
    with pytest.raises(WorkbookError):
        store.append_row("conversations", {"not_a_column": 1})


def test_update_row(store):
    task = first(store.rows("teacher_tasks"))
    updated = store.update_row("teacher_tasks", task["id"], {"status": "done"})
    assert updated["status"] == "done"
    assert store.get("teacher_tasks", task["id"])["status"] == "done"
    assert store.update_row("teacher_tasks", "no-such-id", {"status": "done"}) is None


def test_external_edit_is_picked_up_without_restart(store, workbook):
    student = first(store.rows("students"))
    wb = load_workbook(workbook)
    ws = wb["students"]
    headers = [c.value for c in ws[1]]
    ws.cell(2, headers.index("name") + 1).value = "Renamed In Excel"
    wb.save(workbook)
    wb.close()

    assert store.get("students", student["id"])["name"] == "Renamed In Excel"


def test_locked_file_raises_friendly_error(store, monkeypatch):
    def locked(*_args, **_kwargs):
        raise PermissionError("file is open in another program")

    monkeypatch.setattr(excel_store.os, "replace", locked)
    task = first(store.rows("teacher_tasks"))
    with pytest.raises(WorkbookLockedError, match="close it in Excel"):
        store.update_row("teacher_tasks", task["id"], {"status": "done"})
    assert not list(store.path.parent.glob(".tmp-*"))


def test_tenant_and_guardian_lookups(store):
    tenant = first(store.rows("tenants"))
    assert store.default_tenant_id() == tenant["id"]
    assert store.tenant_for_phone_number_id(f"+{tenant['whatsapp_phone_number_id']}")["id"] == tenant["id"]
    assert store.tenant_for_phone_number_id("999") is None

    guardian = first(store.rows("guardians"))
    number = guardian["whatsapp_number"]
    spaced = f"+{number[:2]} {number[2:7]} {number[7:]}"
    assert store.guardian_by_whatsapp(spaced, tenant["id"])["id"] == guardian["id"]
    assert store.guardian_by_whatsapp("910000000000", tenant["id"]) is None


def test_consent_comes_from_consent_given_at(store):
    guardians = store.rows("guardians")
    assert any(store.has_consent(g) for g in guardians)
    assert any(not store.has_consent(g) for g in guardians)
