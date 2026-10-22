"""JSONL export of saved parent–teacher turns. Uses the temp workbook copy."""

import importlib.util
import json
import os
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def load_exporter():
    path = ROOT / "scripts" / "export_parent_teacher_turns.py"
    spec = importlib.util.spec_from_file_location("export_parent_teacher_turns", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _conversation(store, guardian_id, tenant_id=None, channel="app"):
    return store.append_row("conversations", {
        "tenant_id": tenant_id or store.default_tenant_id(),
        "guardian_id": guardian_id,
        "channel": channel,
        "created_at": "2026-10-01T00:00:00",
    })


def _message(store, conversation_id, role, content, created_at, agent=None, **meta):
    row = {
        "conversation_id": conversation_id,
        "role": role,
        "content": content,
        "created_at": created_at,
        "agent": agent,
        "student_id": meta.get("student_id"),
        "intent_source": meta.get("intent_source"),
        "intent_confidence": meta.get("intent_confidence"),
    }
    return store.append_row("messages", row)


def _seed_pair(store, conversation_id, question, answer, created_at, agent="parent_teacher", **meta):
    _message(store, conversation_id, "user", question, created_at, **meta)
    _message(store, conversation_id, "assistant", answer, created_at, agent=agent, **meta)


def _read(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8")
    if not text.strip():
        return []
    return [json.loads(line) for line in text.splitlines() if line.strip()]


def test_pairs_question_with_reply_in_the_same_conversation(store, tmp_path):
    exporter = load_exporter()
    first = _conversation(store, "g-3-1")
    second = _conversation(store, "g-5-1", channel="whatsapp")
    _seed_pair(store, first["id"], "Who teaches Maths?", "Kavita Nair teaches Mathematics.", "2026-10-06T10:00:00", student_id="s-3-1", intent_source="classifier", intent_confidence=0.91)
    _seed_pair(store, second["id"], "He will be late.", "I will tell the class teacher.", "2026-10-06T11:00:00", student_id="s-5-1", intent_source="llm")
    _message(store, first["id"], "assistant", "Wrong conversation reply.", "2026-10-06T12:00:00", agent="parent_teacher")

    out = tmp_path / "turns.jsonl"
    assert exporter.export_turns(output=out) == 0
    rows = _read(out)
    assert [(row["input"], row["reply"]) for row in rows] == [
        ("Who teaches Maths?", "Kavita Nair teaches Mathematics."),
        ("He will be late.", "I will tell the class teacher."),
    ]
    assert all(row["route"] == "parent_teacher" for row in rows)


def test_default_export_keeps_only_parent_teacher(store, tmp_path):
    exporter = load_exporter()
    conversation = _conversation(store, "g-3-1")
    _seed_pair(store, conversation["id"], "Homework?", "Maths worksheet.", "2026-10-06T09:00:00", agent="daily_update", student_id="s-3-1")
    _seed_pair(store, conversation["id"], "Who teaches Maths?", "Kavita Nair.", "2026-10-06T10:00:00", student_id="s-3-1")

    rows = _read_export(exporter, store, tmp_path)
    assert [row["input"] for row in rows] == ["Who teaches Maths?"]
    assert rows[0]["route"] == "parent_teacher"


def test_multiple_turns_stay_in_order_without_duplicates(store, tmp_path):
    exporter = load_exporter()
    conversation = _conversation(store, "g-3-1")
    _seed_pair(store, conversation["id"], "First question", "First answer", "2026-10-06T10:00:00", student_id="s-3-1")
    _seed_pair(store, conversation["id"], "Second question", "Second answer", "2026-10-06T10:05:00", student_id="s-3-1")

    rows = _read_export(exporter, store, tmp_path)
    assert [(row["input"], row["reply"]) for row in rows] == [
        ("First question", "First answer"),
        ("Second question", "Second answer"),
    ]


def test_since_is_inclusive(store, tmp_path):
    exporter = load_exporter()
    conversation = _conversation(store, "g-3-1")
    _seed_pair(store, conversation["id"], "Before", "Too early", "2026-10-05T23:00:00", student_id="s-3-1")
    _seed_pair(store, conversation["id"], "On the day", "Included", "2026-10-06T00:01:00", student_id="s-3-1")
    _seed_pair(store, conversation["id"], "After", "Included too", "2026-10-07T08:00:00", student_id="s-3-1")

    out = tmp_path / "since.jsonl"
    assert exporter.export_turns(since="2026-10-06", output=out) == 0
    assert [row["input"] for row in _read(out)] == ["On the day", "After"]


def test_agent_flag_exports_that_route_only(store, tmp_path):
    exporter = load_exporter()
    conversation = _conversation(store, "g-3-1")
    _seed_pair(store, conversation["id"], "Who teaches Maths?", "Kavita Nair.", "2026-10-06T10:00:00")
    _seed_pair(store, conversation["id"], "Homework?", "Worksheet.", "2026-10-06T11:00:00", agent="daily_update", student_id="s-3-1")

    out = tmp_path / "daily.jsonl"
    assert exporter.export_turns(agent="daily_update", output=out) == 0
    rows = _read(out)
    assert [row["input"] for row in rows] == ["Homework?"]
    assert rows[0]["route"] == "daily_update"


def test_privacy_omits_contact_columns_and_redacts_numbers_in_text(store, tmp_path):
    exporter = load_exporter()
    tenant = store.default_tenant_id()
    guardian = store.append_row("guardians", {
        "student_id": "s-3-1",
        "name": "ZZ-GUARDIAN-NAME",
        "relation": "parent",
        "whatsapp_number": "900000000088",
        "phone_number": "900000000099",
        "preferred_language": "en",
        "consent_given_at": "2026-01-01",
    })
    conversation = _conversation(store, guardian["id"], tenant)
    _seed_pair(
        store,
        conversation["id"],
        "Call me on 900000000099 please",
        "I will pass that to the teacher.",
        "2026-10-06T10:00:00",
        student_id="s-3-1",
        intent_source="classifier",
        intent_confidence=0.8,
    )

    blob = (tmp_path / "private.jsonl")
    assert exporter.export_turns(output=blob) == 0
    text = blob.read_text(encoding="utf-8")
    assert "ZZ-GUARDIAN-NAME" not in text
    assert "900000000099" not in text
    assert "900000000088" not in text
    assert "[redacted]" in text
    row = json.loads(text)
    assert set(row) == {
        "input", "reply", "route", "intent_source", "intent_confidence", "student_id", "tenant_id", "created_at",
    }


def test_guardian_name_in_message_text_is_redacted_without_touching_ordinary_words(store, tmp_path, capsys, caplog):
    exporter = load_exporter()
    caplog.set_level(0)
    guardian = store.append_row("guardians", {
        "student_id": "s-3-1",
        "name": "Will ZZGuardianname",
        "relation": "parent",
        "whatsapp_number": "900000000088",
        "phone_number": "900000000099",
        "preferred_language": "en",
        "consent_given_at": "2026-01-01",
    })
    conversation = _conversation(store, guardian["id"])
    _seed_pair(
        store,
        conversation["id"],
        "Will ZZGuardianname asked about homework. Call 900000000099. ZZGuardiannameless stays.",
        "The teacher will share today's homework. WhatsApp +91 90000 00088 was not used.",
        "2026-10-06T10:00:00",
        student_id="s-3-1",
        intent_source="classifier",
        intent_confidence=0.8,
    )

    out = tmp_path / "named.jsonl"
    assert exporter.export_turns(output=out) == 0
    text = out.read_text(encoding="utf-8")
    row = json.loads(text)
    for secret in ("900000000099", "900000000088", "+91 90000 00088", "Will ZZGuardianname"):
        assert secret not in text
    without_decoy = text.replace("ZZGuardiannameless", "")
    assert "ZZGuardianname" not in without_decoy
    assert "homework" in row["input"] and "homework" in row["reply"]
    assert "ZZGuardiannameless stays." in row["input"]
    assert row["reply"] == "The teacher will share today's homework. WhatsApp [redacted] was not used."
    captured = capsys.readouterr()
    logged = caplog.text + captured.out + captured.err
    for secret in ("ZZGuardianname", "900000000099", "900000000088", "Will ZZGuardianname"):
        assert secret not in logged


def test_phone_redaction_keeps_iso_dates():
    exporter = load_exporter()
    assert exporter._redact("2026-10-06", []) == "2026-10-06"
    assert exporter._redact("call 2026-10-06 please", []) == "call 2026-10-06 please"
    assert exporter._redact("+91 90000 00088", []) == "[redacted]"
    assert exporter._redact("900000000099", []) == "[redacted]"
    assert exporter._redact("Call +91-90000-00088 on 2026-10-06.", []) == "Call [redacted] on 2026-10-06."


def test_incomplete_pairs_and_missing_student_are_safe(store, tmp_path):
    exporter = load_exporter()
    conversation = _conversation(store, "g-3-1")
    _message(store, conversation["id"], "assistant", "Orphan reply", "2026-10-06T09:00:00", agent="parent_teacher")
    _seed_pair(store, conversation["id"], "Real question", "Real answer", "2026-10-06T10:02:00", intent_source="llm")
    _message(store, conversation["id"], "user", "Nobody answered", "2026-10-06T10:03:00", student_id="s-3-1")

    rows = _read_export(exporter, store, tmp_path)
    assert [row["input"] for row in rows] == ["Real question"]
    assert rows[0]["student_id"] == ""
    assert rows[0]["intent_confidence"] is None
    assert rows[0]["intent_source"] == "llm"


def test_invalid_cli_arguments_exit_nonzero():
    script = ROOT / "scripts" / "export_parent_teacher_turns.py"
    env = os.environ.copy()
    env["OPENAI_API_KEY"] = "test-key"
    bad_date = subprocess.run(
        [sys.executable, str(script), "--since", "06-10-2026"],
        cwd=ROOT, env=env, capture_output=True, text=True,
    )
    assert bad_date.returncode != 0
    assert "YYYY-MM-DD" in (bad_date.stderr + bad_date.stdout)

    exporter = load_exporter()
    assert exporter.main(["--agent", "   "]) == 2


def test_jsonl_lines_parse_and_have_required_types(store, tmp_path):
    exporter = load_exporter()
    conversation = _conversation(store, "g-3-1")
    _seed_pair(
        store, conversation["id"], "Question", "Answer", "2026-10-06T10:15:00",
        student_id="s-3-1", intent_source="classifier", intent_confidence=0.96,
    )
    out = tmp_path / "typed.jsonl"
    exporter.export_turns(output=out)
    line = out.read_text(encoding="utf-8").splitlines()[0]
    row = json.loads(line)
    assert isinstance(row["intent_confidence"], float)
    assert row["intent_confidence"] == 0.96
    for field in ("input", "reply", "route", "intent_source", "student_id", "tenant_id", "created_at"):
        assert isinstance(row[field], str) and row[field]


def test_new_columns_are_added_once_and_old_rows_remain(store, workbook):
    legacy = store.append_row("messages", {
        "conversation_id": "c-legacy",
        "role": "user",
        "content": "keep me",
        "created_at": "2026-09-01T00:00:00",
    })
    wb = load_workbook(workbook)
    ws = wb["messages"]
    headers = [cell.value for cell in ws[1]]
    for name in ("student_id", "intent_source", "intent_confidence"):
        ws.cell(1, headers.index(name) + 1).value = None
    wb.save(workbook)
    wb.close()

    first = store.ensure_columns("messages", ["student_id", "intent_source", "intent_confidence"])
    second = store.ensure_columns("messages", ["student_id", "intent_source", "intent_confidence"])
    assert first == second
    assert first.count("student_id") == 1
    assert first.count("intent_source") == 1
    assert first.count("intent_confidence") == 1
    kept = store.get("messages", legacy["id"])
    assert kept["content"] == "keep me"
    assert kept["conversation_id"] == "c-legacy"
    assert kept.get("student_id") in (None, "")


def test_export_reads_sheets_through_get_store(store, tmp_path, monkeypatch):
    exporter = load_exporter()
    seen = []

    class Recording:
        def rows(self, sheet):
            seen.append(sheet)
            return store.rows(sheet)

        def __getattr__(self, name):
            return getattr(store, name)

    monkeypatch.setattr(exporter, "get_store", lambda: Recording())
    assert exporter.export_turns(output=tmp_path / "out.jsonl") == 0
    assert "conversations" in seen and "messages" in seen


def test_tenants_and_conversations_are_not_crossed(store, tmp_path):
    exporter = load_exporter()
    school = store.default_tenant_id()
    other = _conversation(store, "g-3-1", tenant_id="other-school")
    home = _conversation(store, "g-5-1", tenant_id=school, channel="whatsapp")
    _message(store, other["id"], "user", "Other school question", "2026-10-06T10:00:00", student_id="s-3-1")
    _message(store, home["id"], "assistant", "Must not answer the other school", "2026-10-06T10:01:00", agent="parent_teacher")
    _seed_pair(store, home["id"], "This school's question", "This school's answer", "2026-10-06T10:02:00", student_id="s-5-1")

    rows = _read_export(exporter, store, tmp_path)
    assert [(row["input"], row["tenant_id"]) for row in rows] == [
        ("This school's question", school),
    ]


def test_default_path_and_gitignore(tmp_path, monkeypatch):
    exporter = load_exporter()
    dest = tmp_path / "parent_teacher_turns.jsonl"
    monkeypatch.setattr(exporter, "DEFAULT_OUTPUT", dest)
    assert exporter.main([]) == 0
    assert dest.is_file()
    assert not (ROOT / "data" / "exports" / "parent_teacher_turns.jsonl").exists()

    ignored = subprocess.run(
        ["git", "check-ignore", "-q", "data/exports/parent_teacher_turns.jsonl"],
        cwd=ROOT,
    )
    assert ignored.returncode == 0


def _read_export(exporter, store, tmp_path):
    out = tmp_path / "turns.jsonl"
    assert exporter.export_turns(output=out) == 0
    return _read(out)
