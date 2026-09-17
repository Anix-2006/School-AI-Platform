"""Teacher workspace data lives in an Excel workbook (no Postgres required)."""

from __future__ import annotations

from datetime import date, datetime, time, timedelta
from pathlib import Path
from threading import Lock

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.config import settings

DEFAULT_TEACHER_ID = "t1"
DAY_NAMES = {1: "Monday", 2: "Tuesday", 3: "Wednesday", 4: "Thursday", 5: "Friday"}
_LOCK = Lock()

SHEETS = (
    "teachers",
    "subjects",
    "assignments",
    "timetable",
    "tasks",
    "syllabus",
    "students",
    "attendance",
)


def workbook_path() -> Path:
    raw = Path(settings.teacher_excel_path)
    if raw.is_absolute():
        return raw
    return Path(__file__).resolve().parents[2] / raw


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return str(value).strip()


def _rows(ws: Worksheet) -> list[dict]:
    values = list(ws.iter_rows(values_only=True))
    if not values:
        return []
    headers = [_cell(h) for h in values[0]]
    out = []
    for raw in values[1:]:
        if raw is None or all(v is None or str(v).strip() == "" for v in raw):
            continue
        out.append({headers[i]: _cell(raw[i]) if i < len(raw) else "" for i in range(len(headers))})
    return out


def _parse_date(value: str) -> date | None:
    if not value:
        return None
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    return date.fromisoformat(value[:10])


def _parse_time(value: str) -> time | None:
    if not value:
        return None
    parts = value.replace(".", ":").split(":")
    return time(int(parts[0]), int(parts[1]) if len(parts) > 1 else 0)


def _parse_bool(value: str) -> bool | None:
    if value == "":
        return None
    return value.lower() in ("1", "true", "yes", "y")


def write_demo_workbook(path: Path | None = None) -> Path:
    path = path or workbook_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    today = date.today()
    tenant = settings.default_tenant_id
    wb = Workbook()

    sheets = {
        "teachers": [
            ["id", "tenant_id", "name", "email", "role", "active"],
            ["t1", tenant, "Priya Sharma", "priya.sharma@demo-school.edu", "class_teacher", "true"],
            ["t2", tenant, "Arjun Rao", "arjun.rao@demo-school.edu", "subject_teacher", "true"],
        ],
        "subjects": [
            ["id", "teacher_id", "subject", "grade", "section"],
            ["ts1", "t1", "Math", "3", "A"],
            ["ts2", "t1", "English", "3", "A"],
            ["ts3", "t2", "Science", "3", "A"],
        ],
        "assignments": [
            ["id", "teacher_id", "student_id", "relationship", "subject"],
            ["ta1", "t1", "s1", "class_teacher", ""],
            ["ta2", "t1", "s2", "class_teacher", ""],
            ["ta3", "t1", "s1", "subject_teacher", "Math"],
            ["ta4", "t1", "s2", "subject_teacher", "Math"],
            ["ta5", "t1", "s1", "subject_teacher", "English"],
            ["ta6", "t1", "s2", "subject_teacher", "English"],
            ["ta7", "t2", "s1", "subject_teacher", "Science"],
            ["ta8", "t2", "s2", "subject_teacher", "Science"],
        ],
        "timetable": [
            ["id", "teacher_id", "day_of_week", "period", "start_time", "end_time", "grade", "section", "subject", "room"],
            ["tt-mon-1", "t1", "1", "1", "08:30", "09:10", "3", "A", "Math", "R12"],
            ["tt-mon-2", "t1", "1", "2", "09:15", "09:55", "3", "A", "English", "R12"],
            ["tt-mon-3", "t1", "1", "3", "10:10", "10:50", "3", "A", "Math", "R12"],
            ["tt-tue-1", "t1", "2", "1", "08:30", "09:10", "3", "A", "English", "R12"],
            ["tt-tue-2", "t1", "2", "2", "09:15", "09:55", "3", "A", "Math", "R12"],
            ["tt-wed-1", "t1", "3", "1", "08:30", "09:10", "3", "A", "Math", "R12"],
            ["tt-wed-2", "t1", "3", "2", "09:15", "09:55", "3", "A", "English", "Lab"],
            ["tt-thu-1", "t1", "4", "1", "08:30", "09:10", "3", "A", "Math", "R12"],
            ["tt-thu-2", "t1", "4", "3", "10:10", "10:50", "3", "A", "English", "R12"],
            ["tt-fri-1", "t1", "5", "1", "08:30", "09:10", "3", "A", "English", "R12"],
            ["tt-fri-2", "t1", "5", "2", "09:15", "09:55", "3", "A", "Math", "R12"],
        ],
        "tasks": [
            ["id", "teacher_id", "title", "due_date", "status", "related_subject"],
            ["tk1", "t1", "Mark FA1 Math worksheets", (today + timedelta(days=1)).isoformat(), "todo", "Math"],
            ["tk2", "t1", "Call Rahul's parent about homework", today.isoformat(), "todo", ""],
            ["tk3", "t1", "Submit yesterday's attendance", (today - timedelta(days=1)).isoformat(), "done", ""],
            ["tk4", "t1", "Prep geometry lesson for Friday", (today + timedelta(days=2)).isoformat(), "todo", "Math"],
            ["tk5", "t1", "Upload class photos to diary", (today - timedelta(days=2)).isoformat(), "done", "English"],
        ],
        "syllabus": [
            ["id", "tenant_id", "grade", "subject", "chapter", "status"],
            ["sy1", tenant, "3", "Math", "Fractions", "completed"],
            ["sy2", tenant, "3", "Math", "Multiplication", "in_progress"],
            ["sy3", tenant, "3", "Math", "Geometry", "not_started"],
            ["sy4", tenant, "3", "English", "Nouns", "completed"],
            ["sy5", tenant, "3", "English", "Story writing", "in_progress"],
        ],
        "students": [
            ["id", "tenant_id", "name", "grade", "section", "age_tier"],
            ["s1", tenant, "Ananya", "3", "A", "primary"],
            ["s2", tenant, "Rahul", "3", "A", "primary"],
            ["s3", tenant, "Meera", "4", "B", "primary"],
        ],
        "attendance": [
            ["id", "student_id", "date", "present"],
            [f"att-s1-{today.isoformat()}", "s1", today.isoformat(), "true"],
            [f"att-s2-{today.isoformat()}", "s2", today.isoformat(), "false"],
        ],
    }

    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        for row in rows:
            ws.append(row)
    wb.save(path)
    return path


def ensure_workbook() -> Path:
    path = workbook_path()
    if not path.is_file():
        write_demo_workbook(path)
    return path


def _load() -> dict[str, list[dict]]:
    path = ensure_workbook()
    with _LOCK:
        wb = load_workbook(path, data_only=True)
        data = {name: _rows(wb[name]) if name in wb.sheetnames else [] for name in SHEETS}
        wb.close()
    return data


def get_teacher(tenant_id: str, teacher_id: str) -> dict | None:
    for row in _load()["teachers"]:
        active = _parse_bool(row.get("active", "true"))
        if row.get("id") == teacher_id and row.get("tenant_id") == tenant_id and active is not False:
            return row
    return None


def teacher_profile(tenant_id: str, teacher_id: str) -> dict | None:
    teacher = get_teacher(tenant_id, teacher_id)
    if not teacher:
        return None
    subjects = [
        {
            "subject": r["subject"],
            "grade": r["grade"],
            "section": r.get("section") or "",
        }
        for r in _load()["subjects"]
        if r.get("teacher_id") == teacher_id
    ]
    return {
        "id": teacher["id"],
        "name": teacher["name"],
        "email": teacher.get("email") or None,
        "role": teacher["role"],
        "subjects": subjects,
    }


def teacher_timetable(tenant_id: str, teacher_id: str) -> list[dict]:
    if not get_teacher(tenant_id, teacher_id):
        return []
    slots = []
    for r in _load()["timetable"]:
        if r.get("teacher_id") != teacher_id:
            continue
        day = int(r["day_of_week"])
        slots.append(
            {
                "id": r["id"],
                "day_of_week": day,
                "day_name": DAY_NAMES.get(day, str(day)),
                "period": int(r["period"]),
                "start_time": _parse_time(r.get("start_time", "")),
                "end_time": _parse_time(r.get("end_time", "")),
                "grade": r["grade"],
                "section": r.get("section") or "",
                "subject": r["subject"],
                "room": r.get("room") or None,
            }
        )
    slots.sort(key=lambda s: (s["day_of_week"], s["period"]))
    return slots


def teacher_tasks(tenant_id: str, teacher_id: str) -> list[dict]:
    if not get_teacher(tenant_id, teacher_id):
        return []
    tasks = []
    for r in _load()["tasks"]:
        if r.get("teacher_id") != teacher_id:
            continue
        tasks.append(
            {
                "id": r["id"],
                "title": r["title"],
                "due_date": _parse_date(r.get("due_date", "")),
                "status": r.get("status") or "todo",
                "related_subject": r.get("related_subject") or None,
            }
        )
    tasks.sort(key=lambda t: (0 if t["status"] == "todo" else 1, t["due_date"] or date.max))
    return tasks


def patch_task(tenant_id: str, teacher_id: str, task_id: str, status: str) -> dict | None:
    if not get_teacher(tenant_id, teacher_id):
        return None
    path = ensure_workbook()
    with _LOCK:
        wb = load_workbook(path)
        ws = wb["tasks"]
        headers = [_cell(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
        id_col = headers.index("id") + 1
        teacher_col = headers.index("teacher_id") + 1
        status_col = headers.index("status") + 1
        updated = None
        for row_idx in range(2, ws.max_row + 1):
            if _cell(ws.cell(row_idx, id_col).value) != task_id:
                continue
            if _cell(ws.cell(row_idx, teacher_col).value) != teacher_id:
                continue
            ws.cell(row_idx, status_col).value = status
            updated = True
            break
        if not updated:
            wb.close()
            return None
        wb.save(path)
        wb.close()
    for task in teacher_tasks(tenant_id, teacher_id):
        if task["id"] == task_id:
            return task
    return None


def teacher_syllabus(tenant_id: str, teacher_id: str) -> list[dict]:
    teacher = get_teacher(tenant_id, teacher_id)
    if not teacher:
        return []
    pairs = {(r["grade"], r["subject"]) for r in _load()["subjects"] if r.get("teacher_id") == teacher_id}
    rows = []
    for r in _load()["syllabus"]:
        if r.get("tenant_id") != tenant_id:
            continue
        if (r["grade"], r["subject"]) not in pairs:
            continue
        rows.append(
            {
                "id": r["id"],
                "grade": r["grade"],
                "subject": r["subject"],
                "chapter": r["chapter"],
                "status": r["status"],
                "planned_date": None,
                "completed_date": None,
            }
        )
    rows.sort(key=lambda r: (r["subject"], r["chapter"]))
    return rows


def teacher_students(tenant_id: str, teacher_id: str) -> list[dict]:
    if not get_teacher(tenant_id, teacher_id):
        return []
    data = _load()
    student_ids = {r["student_id"] for r in data["assignments"] if r.get("teacher_id") == teacher_id}
    today = date.today().isoformat()
    attendance = {
        r["student_id"]: _parse_bool(r.get("present", ""))
        for r in data["attendance"]
        if r.get("date", "")[:10] == today
    }
    students = []
    for r in data["students"]:
        if r.get("id") not in student_ids or r.get("tenant_id") != tenant_id:
            continue
        students.append(
            {
                "id": r["id"],
                "name": r["name"],
                "grade": r["grade"],
                "section": r.get("section") or None,
                "age_tier": r.get("age_tier") or "primary",
                "present_today": attendance.get(r["id"]),
            }
        )
    students.sort(key=lambda s: (s["grade"], s["name"]))
    return students
