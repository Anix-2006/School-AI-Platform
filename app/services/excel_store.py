"""All app data lives in one workbook (settings.school_data_path), one
sheet per table with a header row. Every read and write goes through
get_store(); nothing else opens the file.

Reads are cached and reloaded when the file's modified time changes, so
edits saved in Excel show up on the next request. Writes save to a temp
file and swap it in, so a crash mid-save can't leave a half-written
workbook. Run a single uvicorn worker: the lock below is per-process."""

from __future__ import annotations

import logging
import os
import re
import tempfile
import threading
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

from app.config import settings

logger = logging.getLogger(__name__)

DATE_COLUMNS = {
    "date", "date_of_birth", "consent_given_at", "due_date",
    "planned_date", "completed_date", "start_date", "end_date",
}
DATETIME_COLUMNS = {"created_at"}
TIME_COLUMNS = {"start_time", "end_time"}
BOOL_COLUMNS = {"active", "present"}
NUMBER_COLUMNS = {"marks_obtained", "marks_total", "chapter_no", "intent_confidence"}

PROJECT_ROOT = Path(__file__).resolve().parents[2]


class WorkbookError(RuntimeError):
    """The workbook is missing, unreadable, or doesn't have the expected sheet/columns."""


class WorkbookLockedError(WorkbookError):
    """The workbook couldn't be saved - on Windows this usually means it's open in Excel."""


def _to_number(value):
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else value
    try:
        number = float(str(value))
    except ValueError:
        return None
    return int(number) if number.is_integer() else number


def _clean(column: str, value):
    """Normalise one cell so callers see the same types whether a value was
    typed as text or as a native Excel date/number/boolean."""
    if isinstance(value, str):
        value = value.strip()
    if value is None or value == "":
        return None
    if column in BOOL_COLUMNS:
        if isinstance(value, bool):
            return value
        return str(value).strip().lower() in ("true", "yes", "y", "1")
    if column in DATE_COLUMNS:
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        return str(value)[:10]
    if column in DATETIME_COLUMNS:
        if isinstance(value, datetime):
            return value.isoformat(timespec="seconds")
        return str(value)
    if column in TIME_COLUMNS:
        if isinstance(value, (time, datetime)):
            return value.strftime("%H:%M")
        return str(value)
    if column in NUMBER_COLUMNS:
        return _to_number(value)
    if isinstance(value, bool):
        return value
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def digits(value) -> str:
    return re.sub(r"\D", "", str(value or ""))


class ExcelStore:
    def __init__(self, path: Path):
        self.path = Path(path)
        self._lock = threading.RLock()
        self._mtime: float | None = None
        self._sheets: dict[str, list[dict]] = {}
        self._headers: dict[str, list[str]] = {}

    # ---- loading -------------------------------------------------------

    def _ensure_loaded(self) -> None:
        try:
            mtime = self.path.stat().st_mtime
        except FileNotFoundError:
            raise WorkbookError(f"Workbook not found: {self.path}") from None
        if mtime == self._mtime:
            return
        with self._lock:
            if mtime == self._mtime:
                return
            try:
                wb = load_workbook(self.path, read_only=True, data_only=True)
            except (zipfile.BadZipFile, OSError, KeyError) as exc:
                if self._sheets:
                    logger.warning("Workbook unreadable (%s); keeping the last good copy", exc)
                    return
                raise WorkbookError(f"Couldn't read {self.path}: {exc}") from exc
            sheets, headers = {}, {}
            try:
                for ws in wb.worksheets:
                    rows = ws.iter_rows(values_only=True)
                    header_row = next(rows, None) or ()
                    cols = [str(h).strip() if h is not None else "" for h in header_row]
                    headers[ws.title] = cols
                    records = []
                    for raw in rows:
                        if raw is None or all(v is None or str(v).strip() == "" for v in raw):
                            continue
                        records.append({
                            col: _clean(col, raw[i] if i < len(raw) else None)
                            for i, col in enumerate(cols) if col
                        })
                    sheets[ws.title] = records
            finally:
                wb.close()
            self._sheets, self._headers, self._mtime = sheets, headers, mtime

    # ---- generic reads -------------------------------------------------

    def rows(self, sheet: str) -> list[dict]:
        self._ensure_loaded()
        if sheet not in self._sheets:
            raise WorkbookError(f"Sheet '{sheet}' is missing from {self.path.name}")
        return [dict(r) for r in self._sheets[sheet]]

    def get(self, sheet: str, row_id: str) -> dict | None:
        return next((r for r in self.rows(sheet) if r.get("id") == row_id), None)

    def find(self, sheet: str, **filters) -> list[dict]:
        return [r for r in self.rows(sheet) if all(r.get(k) == v for k, v in filters.items())]

    # ---- writes --------------------------------------------------------

    def append_rows(self, sheet: str, rows: list[dict]) -> list[dict]:
        """Append rows (one save for all of them). Missing ids are generated."""
        if not rows:
            return []
        with self._lock:
            wb = self._open_for_write()
            ws = self._sheet_for_write(wb, sheet)
            headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
            written = []
            for row in rows:
                unknown = set(row) - set(headers)
                if unknown:
                    wb.close()
                    raise WorkbookError(f"Sheet '{sheet}' has no column(s): {', '.join(sorted(unknown))}")
                row = {"id": f"{sheet}-{uuid4().hex[:10]}", **row}
                ws.append([row.get(h) for h in headers])
                written.append({h: _clean(h, row.get(h)) for h in headers if h})
            self._save(wb)
            return written

    def append_row(self, sheet: str, row: dict) -> dict:
        return self.append_rows(sheet, [row])[0]

    def update_row(self, sheet: str, row_id: str, fields: dict) -> dict | None:
        """Set fields on the row with this id. Returns the updated row, or None if no such id."""
        with self._lock:
            wb = self._open_for_write()
            ws = self._sheet_for_write(wb, sheet)
            headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
            unknown = set(fields) - set(headers)
            if unknown:
                wb.close()
                raise WorkbookError(f"Sheet '{sheet}' has no column(s): {', '.join(sorted(unknown))}")
            id_col = headers.index("id") + 1
            for row_idx in range(2, ws.max_row + 1):
                if _clean("id", ws.cell(row_idx, id_col).value) != row_id:
                    continue
                for name, value in fields.items():
                    ws.cell(row_idx, headers.index(name) + 1).value = value
                updated = {h: _clean(h, ws.cell(row_idx, i + 1).value) for i, h in enumerate(headers) if h}
                self._save(wb)
                return updated
            wb.close()
            return None

    def _open_for_write(self):
        if not self.path.is_file():
            raise WorkbookError(f"Workbook not found: {self.path}")
        return load_workbook(self.path)

    def _sheet_for_write(self, wb, sheet: str):
        if sheet not in wb.sheetnames:
            wb.close()
            raise WorkbookError(f"Sheet '{sheet}' is missing from {self.path.name}")
        return wb[sheet]

    def _save(self, wb) -> None:
        fd, tmp = tempfile.mkstemp(dir=self.path.parent, prefix=".tmp-", suffix=".xlsx")
        os.close(fd)
        try:
            wb.save(tmp)
            os.replace(tmp, self.path)
        except OSError as exc:
            Path(tmp).unlink(missing_ok=True)
            raise WorkbookLockedError(
                f"Couldn't save {self.path.name} - close it in Excel and try again ({exc})."
            ) from exc
        finally:
            wb.close()
        self._mtime = None

    # ---- lookups shared by several features -----------------------------

    def default_tenant_id(self) -> str:
        if settings.default_tenant_id:
            return settings.default_tenant_id
        tenants = self.rows("tenants")
        if not tenants:
            raise WorkbookError("The tenants sheet is empty")
        return tenants[0]["id"]

    def tenant_for_phone_number_id(self, phone_number_id: str) -> dict | None:
        wanted = digits(phone_number_id)
        if not wanted:
            return None
        return next(
            (t for t in self.rows("tenants") if digits(t.get("whatsapp_phone_number_id")) == wanted),
            None,
        )

    def student(self, student_id: str, tenant_id: str | None = None) -> dict | None:
        s = self.get("students", student_id)
        if s is None or (tenant_id and s.get("tenant_id") != tenant_id):
            return None
        return s

    def students(self, tenant_id: str, active_only: bool = True) -> list[dict]:
        return [
            s for s in self.rows("students")
            if s.get("tenant_id") == tenant_id and (not active_only or s.get("active") is not False)
        ]

    def guardian(self, guardian_id: str, tenant_id: str | None = None) -> dict | None:
        g = self.get("guardians", guardian_id)
        if g is None:
            return None
        if tenant_id and self.student(g.get("student_id"), tenant_id) is None:
            return None
        return g

    def guardians(self, tenant_id: str) -> list[dict]:
        ids = {s["id"] for s in self.students(tenant_id, active_only=False)}
        return [g for g in self.rows("guardians") if g.get("student_id") in ids]

    def guardian_by_whatsapp(self, number: str, tenant_id: str) -> dict | None:
        wanted = digits(number)
        if not wanted:
            return None
        return next(
            (g for g in self.guardians(tenant_id) if digits(g.get("whatsapp_number")) == wanted),
            None,
        )

    @staticmethod
    def has_consent(guardian: dict) -> bool:
        return bool(guardian.get("consent_given_at"))

    def teacher(self, teacher_id: str, tenant_id: str) -> dict | None:
        t = self.get("teachers", teacher_id)
        if t is None or t.get("tenant_id") != tenant_id or t.get("active") is False:
            return None
        return t

    def teachers(self, tenant_id: str) -> list[dict]:
        return [
            t for t in self.rows("teachers")
            if t.get("tenant_id") == tenant_id and t.get("active") is not False
        ]


def resolve_path(raw: str) -> Path:
    path = Path(raw)
    return path if path.is_absolute() else PROJECT_ROOT / path


_store: ExcelStore | None = None
_store_lock = threading.Lock()


def get_store() -> ExcelStore:
    """The single shared store. Rebuilt only if settings.school_data_path changes."""
    global _store
    path = resolve_path(settings.school_data_path)
    if _store is None or _store.path != path:
        with _store_lock:
            if _store is None or _store.path != path:
                _store = ExcelStore(path)
    return _store
